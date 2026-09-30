import os
from pathlib import Path
from typing import List, Dict, Any

from langchain_text_splitters import RecursiveCharacterTextSplitter
from pdfminer.high_level import extract_text as extract_pdf_text
import docx
import pytesseract
from PIL import Image

class DocumentProcessor:
    def __init__(self, chunk_size: int = 500, chunk_overlap: int = 50):
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
            is_separator_regex=False,
        )

    def process_document(self, file_path: str) -> List[Dict[str, Any]]:
        """
        Reads a document, extracts text based on file type, and returns chunked text with metadata.
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        ext = path.suffix.lower()
        extracted_text = ""

        if ext == '.pdf':
            extracted_text = self._extract_from_pdf(file_path)
        elif ext == '.docx':
            extracted_text = self._extract_from_docx(file_path)
        elif ext in ['.png', '.jpg', '.jpeg']:
            extracted_text = self._extract_from_image(file_path)
        else:
            raise ValueError(f"Unsupported file type: {ext}")

        if not extracted_text.strip():
            print(f"Warning: No text could be extracted from {file_path}")
            return []

        # Split text into chunks
        chunks = self.text_splitter.split_text(extracted_text)
        
        # Add metadata to each chunk
        documents = []
        for i, chunk in enumerate(chunks):
            documents.append({
                "content": chunk,
                "metadata": {
                    "source": path.name,
                    "type": ext,
                    "chunk_id": i
                }
            })
            
        return documents

    def _extract_from_pdf(self, file_path: str) -> str:
        try:
            return extract_pdf_text(file_path)
        except Exception as e:
            print(f"Error extracting PDF text: {e}")
            return ""

    def _extract_from_docx(self, file_path: str) -> str:
        try:
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                full_text.append(para.text)
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    row_data = []
                    for cell in row.cells:
                        row_data.append(cell.text.strip())
                    full_text.append(" | ".join(row_data))
                    
            return '\n'.join(full_text)
        except Exception as e:
            print(f"Error extracting DOCX text: {e}")
            return ""

    def _extract_from_image(self, file_path: str) -> str:
        try:
            image = Image.open(file_path)
            return pytesseract.image_to_string(image)
        except Exception as e:
            print(f"Error extracting Image text: {e}")
            return ""

if __name__ == "__main__":
    # Simple test
    processor = DocumentProcessor()
    print("DocumentProcessor initialized. Ready to process files.")
